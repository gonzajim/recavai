#!/usr/bin/env python3
"""
Inventario del corpus documental: título, tamaño en palabras y tamaño de fichero.

Los ficheros del corpus no están en el repositorio —solo sus vectores, en Pinecone—,
así que hay que apuntar a la carpeta que los contiene.

Con --pinecone, además, contrasta la carpeta con lo que está REALMENTE indexado. Esa
comprobación es la que importa: el sistema responde con lo que hay en el índice, no
con lo que haya en una carpeta del portátil, y los dos pueden haberse desincronizado.

Uso:
  python scripts/corpus_inventory.py --corpus "/ruta/al/corpus"
  python scripts/corpus_inventory.py --corpus "/ruta/al/corpus" --pinecone
  python scripts/corpus_inventory.py --corpus "/ruta" --out results/inventario
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, asdict
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}
_WORD = re.compile(r"\S+")


@dataclass
class Doc:
    source: str          # ruta relativa sin extensión — la misma clave que usa Pinecone
    titulo: str          # nombre de fichero legible: es lo que el sistema cita
    titulo_interno: str  # primer titular dentro del documento (solo informativo)
    carpeta: str
    formato: str
    paginas: int
    palabras: int
    bytes: int
    vectores: int = 0          # solo con --pinecone
    palabras_indexadas: int = 0  # palabras realmente guardadas en el índice
    error: str = ""

    @property
    def cobertura(self) -> float:
        """Qué fracción del texto del documento llegó al índice. Por debajo de 1 significa
        que hay contenido del PDF que el asistente NO puede recuperar ni citar."""
        return (self.palabras_indexadas / self.palabras) if self.palabras else 0.0

    @property
    def mb(self) -> float:
        return self.bytes / (1024 * 1024)


def human(n: int) -> str:
    for unit in ("B", "KB", "MB"):
        if n < 1024 or unit == "MB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} MB"


# ---------------------------------------------------------------------------
# Extracción de texto
# ---------------------------------------------------------------------------
def read_pdf(path: Path) -> tuple[str, int, str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", 0, "falta pypdf"
    try:
        reader = PdfReader(str(path))
        pages = len(reader.pages)
        chunks = []
        for p in reader.pages:
            try:
                chunks.append(p.extract_text() or "")
            except Exception:                       # noqa: BLE001
                pass            # una página ilegible no invalida el resto del documento
        return "\n".join(chunks), pages, ""
    except Exception as e:                          # noqa: BLE001
        return "", 0, f"{type(e).__name__}"


def read_docx(path: Path) -> tuple[str, int, str]:
    try:
        import docx
    except ImportError:
        return "", 0, "falta python-docx"
    try:
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts), 0, ""              # DOCX no tiene paginación fija
    except Exception as e:                          # noqa: BLE001
        return "", 0, f"{type(e).__name__}"


def read_plain(path: Path) -> tuple[str, int, str]:
    try:
        return path.read_text(encoding="utf-8", errors="replace"), 0, ""
    except Exception as e:                          # noqa: BLE001
        return "", 0, f"{type(e).__name__}"


def title_of(path: Path) -> str:
    """Título legible a partir del NOMBRE DE FICHERO, no de la primera línea del texto.

    Se intentó lo segundo y salía mal: los PDF de EUR-Lex empiezan con el descargo
    'Este texto es exclusivamente un instrumento de documentación…' y varias guías de
    la OCDE arrancan con un fragmento suelto del título maquetado en varias líneas.
    Además, el nombre de fichero es lo que el sistema cita y lo que aparece en la
    interfaz, así que es el identificador que el lector reconoce.
    """
    name = re.sub(r"^\d+(_\d+)*[_\-\s]+", "", path.stem)   # prefijos de ordenación
    name = name.replace("_", " ").strip()
    return (name or path.stem)[:110]


def inner_title(text: str) -> str:
    """Primer titular con sustancia dentro del documento. Solo informativo: va al CSV."""
    for line in text.splitlines()[:40]:
        s = " ".join(line.split())
        if len(s) > 18 and not s.isdigit() and "instrumento de documentación" not in s:
            return s[:140]
    return ""


def scan(root: Path) -> list[Doc]:
    docs: list[Doc] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED:
            continue
        if path.name.startswith("~$") or path.name == ".DS_Store":
            continue
        ext = path.suffix.lower()
        text, pages, err = (
            read_pdf(path) if ext == ".pdf" else
            read_docx(path) if ext == ".docx" else
            read_plain(path)
        )
        rel = path.relative_to(root)
        docs.append(Doc(
            source=str(rel.with_suffix("")).replace("\\", "/"),
            titulo=title_of(path),
            titulo_interno=inner_title(text),
            carpeta=str(rel.parent) if str(rel.parent) != "." else "(raíz)",
            formato=ext.lstrip("."),
            paginas=pages,
            palabras=len(_WORD.findall(text)),
            bytes=path.stat().st_size,
            error=err,
        ))
    return docs


# ---------------------------------------------------------------------------
# Contraste con lo indexado
# ---------------------------------------------------------------------------
LENGTHS: list[int] = []   # palabras por vector, para la distribución del informe


def pinecone_sources() -> tuple[dict[str, int], dict[str, int]] | None:
    """({source: nº de vectores}, {source: palabras indexadas}) del índice de producción.
    Rellena además LENGTHS con el tamaño de cada vector. None si no se puede consultar."""
    try:
        from pinecone import Pinecone
    except ImportError:
        print("  (falta el paquete pinecone; se omite el contraste)", file=sys.stderr)
        return None
    key, name = os.getenv("PINECONE_API_KEY"), os.getenv("PINECONE_INDEX_NAME")
    if not key or not name:
        print("  (faltan PINECONE_API_KEY / PINECONE_INDEX_NAME; se omite)", file=sys.stderr)
        return None
    try:
        pc = Pinecone(api_key=key)
        index = pc.Index(host=name) if name.startswith("http") else pc.Index(name)

        # El corpus vive en un namespace llamado literalmente "__default__", no en "".
        # Y list() devuelve objetos ListItem, no cadenas: pasarlos tal cual a fetch()
        # no da error, simplemente devuelve cero vectores en silencio.
        stats = index.describe_index_stats()
        spaces = {k: v.get("vector_count", 0) for k, v in (stats.get("namespaces") or {}).items()}
        ns = max(spaces, key=spaces.get) if spaces else ""
        total = spaces.get(ns, 0)

        counts: dict[str, int] = defaultdict(int)
        words: dict[str, int] = defaultdict(int)
        seen = 0
        for page in index.list(namespace=ns):
            ids = [getattr(i, "id", i) for i in page]
            if not ids:
                continue
            vectors = index.fetch(ids=ids, namespace=ns).vectors or {}
            for v in vectors.values():
                md = getattr(v, "metadata", None) or {}
                src = str(md.get("source", "(sin source)"))
                counts[src] += 1
                n = len(_WORD.findall(str(md.get("text", ""))))
                words[src] += n
                LENGTHS.append(n)
                seen += 1
            print(f"\r  leyendo el índice '{ns}'… {seen}/{total} vectores",
                  end="", file=sys.stderr)
        print(file=sys.stderr)
        if seen == 0:
            print("  (el índice devolvió 0 vectores; no se contrasta)", file=sys.stderr)
            return None
        return dict(counts), dict(words)
    except Exception as e:                          # noqa: BLE001
        print(f"  (no se pudo leer Pinecone: {type(e).__name__}: {e})", file=sys.stderr)
        return None


def norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", s.lower())


# ---------------------------------------------------------------------------
def report(docs: list[Doc], indexed: dict[str, int] | None) -> str:
    L: list[str] = []
    total_w = sum(d.palabras for d in docs)
    total_b = sum(d.bytes for d in docs)
    total_p = sum(d.paginas for d in docs)

    L.append("# Corpus documental de RecavAI\n")
    L.append(f"**{len(docs)} documentos · {total_w:,} palabras · {total_p:,} páginas · "
             f"{human(total_b)}**\n".replace(",", "."))

    by_folder: dict[str, list[Doc]] = defaultdict(list)
    for d in docs:
        by_folder[d.carpeta].append(d)

    L.append("## Resumen por carpeta\n")
    L.append("| Carpeta | Docs | Palabras | Páginas | Tamaño |")
    L.append("|---|---:|---:|---:|---:|")
    for folder in sorted(by_folder):
        g = by_folder[folder]
        L.append(f"| {folder} | {len(g)} | {sum(x.palabras for x in g):,} | "
                 f"{sum(x.paginas for x in g):,} | {human(sum(x.bytes for x in g))} |"
                 .replace(",", "."))
    L.append(f"| **Total** | **{len(docs)}** | **{total_w:,}** | **{total_p:,}** | "
             f"**{human(total_b)}** |".replace(",", "."))

    L.append("\n## Documentos\n")
    head = "| # | Documento | Carpeta | Fmt | Págs | Palabras | Tamaño |"
    sep = "|---:|---|---|---|---:|---:|---:|"
    if indexed is not None:
        head += " Vectores | Palabras indexadas | Cobertura |"
        sep += "---:|---:|---:|"
    L.append(head)
    L.append(sep)
    for i, d in enumerate(sorted(docs, key=lambda x: (-x.palabras)), 1):
        title = d.titulo if not d.error else f"{d.titulo} ⚠️ {d.error}"
        row = (f"| {i} | {title} | {d.carpeta} | {d.formato} | "
               f"{d.paginas or '—'} | {d.palabras:,} | {human(d.bytes)} |".replace(",", "."))
        if indexed is not None:
            row += (f" {d.vectores or '—'} | {d.palabras_indexadas:,} | "
                    f"{d.cobertura:.0%} |").replace(",", ".")
        L.append(row)

    fallidos = [d for d in docs if d.error or d.palabras == 0]
    if fallidos:
        L.append("\n## Documentos de los que no se extrajo texto\n")
        L.append("Probablemente escaneados sin capa de texto: ocupan sitio en el corpus "
                 "pero **no aportan nada recuperable** salvo que se les pase OCR.\n")
        for d in fallidos:
            L.append(f"- `{d.source}` ({human(d.bytes)}"
                     + (f", {d.paginas} págs" if d.paginas else "")
                     + (f", {d.error}" if d.error else "") + ")")

    if indexed is not None:
        sin_indexar = [d for d in docs if d.vectores == 0]
        local = {norm(d.source) for d in docs}
        huerfanos = [s for s in indexed if norm(s) not in local
                     and not any(norm(s) in norm(d.source) or norm(d.source) in norm(s) for d in docs)]
        idx_w = sum(d.palabras_indexadas for d in docs)
        src_w = sum(d.palabras for d in docs)
        L.append("\n## Contraste con el índice\n")
        L.append(f"- Vectores en el índice: **{sum(indexed.values()):,}**".replace(",", "."))
        L.append(f"- Palabras del corpus: **{src_w:,}**".replace(",", "."))
        L.append(f"- Palabras realmente indexadas: **{idx_w:,}** "
                 f"(**{idx_w/src_w:.1%}** del corpus)".replace(",", "."))
        L.append(f"- Media de palabras por vector: **{idx_w/max(1,sum(indexed.values())):.1f}**")
        if LENGTHS:
            xs = sorted(LENGTHS)
            def pc(q): return xs[min(len(xs) - 1, int(len(xs) * q))]
            tiny = sum(1 for x in xs if x <= 5)
            big = [x for x in xs if x >= 50]
            L.append("\n### Distribución del tamaño de los fragmentos\n")
            L.append("Importa porque un índice troceado en fragmentos diminutos recupera mal: "
                     "el fragmento no contiene la respuesta entera.\n")
            L.append("| p10 | p25 | mediana | p75 | p90 | p99 | máx |")
            L.append("|---:|---:|---:|---:|---:|---:|---:|")
            L.append(f"| {pc(.10)} | {pc(.25)} | {pc(.50)} | {pc(.75)} | "
                     f"{pc(.90)} | {pc(.99)} | {xs[-1]} |")
            L.append(f"\n- Fragmentos de 5 palabras o menos: **{tiny}** "
                     f"({tiny/len(xs):.0%}), que contienen el "
                     f"**{sum(x for x in xs if x <= 5)/max(1,sum(xs)):.1%}** del texto.")
            L.append(f"- Fragmentos de 50 palabras o más: **{len(big)}** "
                     f"({len(big)/len(xs):.0%}), que contienen el "
                     f"**{sum(big)/max(1,sum(xs)):.1%}** del texto.")
        pobres = sorted((d for d in docs if d.vectores and d.cobertura < 0.5),
                        key=lambda x: x.cobertura)
        if pobres:
            L.append(f"\n### Documentos con menos de la mitad del texto indexado "
                     f"({len(pobres)} de {len([d for d in docs if d.vectores])})\n")
            L.append("El asistente no puede citar lo que no está en el índice.\n")
            L.append("| Documento | Palabras | Indexadas | Cobertura |")
            L.append("|---|---:|---:|---:|")
            for d in pobres[:25]:
                L.append(f"| {d.titulo} | {d.palabras:,} | {d.palabras_indexadas:,} | "
                         f"{d.cobertura:.0%} |".replace(",", "."))
        L.append(f"- Documentos distintos indexados: **{len(indexed)}**")
        L.append(f"- Documentos de la carpeta sin vectores: **{len(sin_indexar)}**")
        if sin_indexar:
            L.append("\n  No están en el índice, así que el asistente NO puede citarlos:\n")
            for d in sin_indexar[:40]:
                L.append(f"  - `{d.source}`")
        if huerfanos:
            L.append(f"\n- Indexados que no están en la carpeta: **{len(huerfanos)}**\n")
            for s in sorted(huerfanos)[:40]:
                L.append(f"  - `{s}` ({indexed[s]} vectores)")

    L.append("\n## Mantenimiento\n")
    L.append("Este inventario dice qué hay, no si sigue siendo válido. Dos comprobaciones "
             "complementarias, cada una con su script:\n")
    L.append("- **¿La normativa sigue vigente?** Los textos de EUR-Lex se consolidan cada vez "
             "que entra en vigor una modificación; un PDF descargado hace meses puede citar "
             "una versión ya superada sin avisar de ello.\n"
             "  ```\n  python scripts/check_eurlex_versions.py --corpus \"<carpeta>\"\n  ```")
    L.append("- **¿El texto indexado está limpio?** Los PDF oficiales parten palabras al final "
             "de línea con un guión blando, y la partición sobrevive a la extracción.\n"
             "  ```\n  python scripts/fix_soft_hyphens.py            # diagnóstico\n"
             "  python scripts/fix_soft_hyphens.py --apply    # repara el índice\n  ```")
    return "\n".join(L)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--corpus", required=True, help="carpeta con los documentos")
    ap.add_argument("--pinecone", action="store_true",
                    help="contrastar con el índice (necesita PINECONE_API_KEY)")
    ap.add_argument("--out", help="prefijo de salida: genera .md y .csv")
    a = ap.parse_args()

    root = Path(a.corpus).expanduser()
    if not root.is_dir():
        sys.exit(f"No existe la carpeta: {root}")

    print(f"Leyendo {root}…", file=sys.stderr)
    docs = scan(root)
    if not docs:
        sys.exit("No se encontró ningún documento soportado (.pdf .docx .txt .md)")

    indexed = None
    if a.pinecone:
        print("Contrastando con el índice de Pinecone…", file=sys.stderr)
        result = pinecone_sources()
        if result is not None:
            indexed, idx_words = result
            for d in docs:
                key = d.source if d.source in indexed else next(
                    (s for s in indexed if norm(s) == norm(d.source)), None)
                if key:
                    d.vectores = indexed[key]
                    d.palabras_indexadas = idx_words.get(key, 0)

    md = report(docs, indexed)
    print(md)

    if a.out:
        out = Path(a.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.with_suffix(".md").write_text(md + "\n", encoding="utf-8")
        with open(out.with_suffix(".csv"), "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(asdict(docs[0])))
            w.writeheader()
            w.writerows(asdict(d) for d in docs)
        print(f"\nEscritos {out.with_suffix('.md')} y {out.with_suffix('.csv')}",
              file=sys.stderr)


if __name__ == "__main__":
    main()
